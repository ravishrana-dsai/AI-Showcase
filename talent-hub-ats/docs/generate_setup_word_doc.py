#!/usr/bin/env python3
"""
Generate SETUP_IDE_PYTHON.docx from the setup guide content.
Run: python3 generate_setup_word_doc.py
Requires: pip install python-docx
Output: SETUP_IDE_PYTHON.docx (in the same folder as this script)
"""

try:
    from docx import Document
    from docx.shared import Pt, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    print("Missing python-docx. Install it with: pip3 install python-docx")
    exit(1)

def add_heading(doc, text, level=1):
    doc.add_heading(text, level=level)

def add_para(doc, text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    if bold:
        run.bold = True
    return p

def add_code(doc, text):
    p = doc.add_paragraph()
    p.style = "Normal"
    run = p.add_run(text)
    run.font.name = "Menlo"
    run.font.size = Pt(10)
    return p

def main():
    doc = Document()
    doc.add_heading("Setting Up the IDE, Python, and Libraries on a New Mac", 0)
    add_para(doc, "Step-by-step guide for a non-technical person. Only covers: Cursor (IDE), Python, and Python libraries. Nothing about project folders or the app itself.")
    doc.add_paragraph()

    # --- Part 1 ---
    add_heading(doc, "Part 1: Install Python", 1)
    add_heading(doc, "Step 1.1: Download Python", 2)
    doc.add_paragraph("1. Open your browser (Safari, Chrome, etc).")
    doc.add_paragraph("2. Go to: https://www.python.org/downloads")
    doc.add_paragraph('3. Click the big yellow button that says "Download Python 3.x.x" (the latest version).')
    doc.add_paragraph("4. The installer file (e.g. python-3.12.x-macos11.pkg) will download to your Downloads folder.")
    doc.add_paragraph()

    add_heading(doc, "Step 1.2: Run the installer", 2)
    doc.add_paragraph("1. Open Finder → Downloads.")
    doc.add_paragraph("2. Double-click the Python installer file (the one you just downloaded).")
    doc.add_paragraph("3. Click Continue through the screens.")
    doc.add_paragraph("4. When asked, read and accept the license, then click Continue and Agree.")
    doc.add_paragraph("5. Leave the default options as they are (just click Continue).")
    doc.add_paragraph("6. Click Install. Enter your Mac password if prompted.")
    doc.add_paragraph('7. Wait until you see "The installation was successful."')
    doc.add_paragraph("8. Click Close.")
    doc.add_paragraph()

    add_heading(doc, "Step 1.3: Check that Python is installed", 2)
    doc.add_paragraph("1. Open Terminal: press Command + Space, type Terminal, press Enter.")
    doc.add_paragraph("2. Type exactly and press Enter:")
    add_code(doc, "   python3 --version")
    doc.add_paragraph("3. You should see something like Python 3.12.0 (the numbers may differ).")
    add_para(doc, 'If you see "command not found", Python did not install correctly. Run the installer again and make sure it finishes without errors.')
    doc.add_paragraph()

    add_heading(doc, "Step 1.4: Check that pip is installed", 2)
    doc.add_paragraph("1. In Terminal, type and press Enter:")
    add_code(doc, "   pip3 --version")
    doc.add_paragraph("2. You should see a line like pip 23.x.x from ...")
    doc.add_paragraph("If pip3 is not found, try:")
    add_code(doc, "   python3 -m pip --version")
    doc.add_paragraph("If that works, you can use python3 -m pip instead of pip3 in the steps below.")
    doc.add_paragraph()

    # --- Part 2 ---
    add_heading(doc, "Part 2: Install Cursor (the IDE)", 1)
    add_heading(doc, "Step 2.1: Download Cursor", 2)
    doc.add_paragraph("1. In your browser, go to: https://cursor.com")
    doc.add_paragraph("2. Click the button to download Cursor for Mac.")
    doc.add_paragraph("3. The file (e.g. Cursor-0.xx.x.dmg) will download to Downloads.")
    doc.add_paragraph()

    add_heading(doc, "Step 2.2: Install Cursor", 2)
    doc.add_paragraph("1. Open Finder → Downloads.")
    doc.add_paragraph("2. Double-click the Cursor .dmg file.")
    doc.add_paragraph("3. A window will open with the Cursor icon. Drag the Cursor icon into the Applications folder (or follow the on-screen instructions).")
    doc.add_paragraph("4. When the copy is done, you can close the window. If you see a disk on your desktop, you can right-click it and choose Eject.")
    doc.add_paragraph("5. Open Cursor from Applications (or press Command + Space, type Cursor, press Enter).")
    doc.add_paragraph()

    add_heading(doc, "Step 2.3: First-time setup (optional)", 2)
    doc.add_paragraph("1. The first time you open Cursor, you may be asked to sign in or create an account. You can do that if you want cloud features.")
    doc.add_paragraph("2. You can skip or complete the welcome / Get Started steps. No special settings are required for basic use.")
    doc.add_paragraph()

    # --- Part 3 ---
    add_heading(doc, "Part 3: Install Python libraries (using pip)", 1)
    add_para(doc, "You install Python libraries from Terminal with pip3 (or python3 -m pip). Run these from any folder; they install for your user account.")
    doc.add_paragraph()

    add_heading(doc, "Step 3.1: Open Terminal", 2)
    doc.add_paragraph("Press Command + Space, type Terminal, press Enter.")
    doc.add_paragraph()

    add_heading(doc, "Step 3.2: Upgrade pip (recommended)", 2)
    doc.add_paragraph("Run this first so you have the latest pip:")
    add_code(doc, "pip3 install --upgrade pip")
    doc.add_paragraph("(If that fails, try: python3 -m pip install --upgrade pip)")
    doc.add_paragraph()

    add_heading(doc, "Step 3.3: Install common libraries", 2)
    doc.add_paragraph("Run each line below one at a time, wait for it to finish, then run the next.")
    add_para(doc, "Data and science:", bold=True)
    add_code(doc, "pip3 install numpy")
    add_code(doc, "pip3 install pandas")
    add_code(doc, "pip3 install matplotlib")
    add_code(doc, "pip3 install scipy")
    add_para(doc, "Jupyter (notebooks):", bold=True)
    add_code(doc, "pip3 install jupyter")
    add_para(doc, "Requests (for HTTP / APIs):", bold=True)
    add_code(doc, "pip3 install requests")
    add_para(doc, "Virtual environments (optional):", bold=True)
    add_code(doc, "pip3 install virtualenv")
    doc.add_paragraph()

    add_heading(doc, "Step 3.4: Check that libraries are installed", 2)
    doc.add_paragraph("1. In Terminal, type:")
    add_code(doc, "pip3 list")
    doc.add_paragraph("2. You should see a long list that includes numpy, pandas, matplotlib, jupyter, requests, etc.")
    doc.add_paragraph("To check one library by name, for example:")
    add_code(doc, 'python3 -c "import numpy; print(numpy.__version__)"')
    doc.add_paragraph("You should see a version number (e.g. 1.26.0).")
    doc.add_paragraph()

    # --- Part 4 ---
    add_heading(doc, "Part 4: Use Python from Cursor", 1)
    doc.add_paragraph("1. Open Cursor.")
    doc.add_paragraph("2. Create or open any .py file (e.g. File → New File, then File → Save As and name it test.py).")
    doc.add_paragraph("3. Type something like:")
    add_code(doc, 'print("Hello")\nimport numpy as np\nprint(np.__version__)')
    doc.add_paragraph("4. Open the Terminal inside Cursor: Terminal → New Terminal (or View → Terminal).")
    doc.add_paragraph("5. In that terminal, run your script:")
    add_code(doc, "python3 test.py")
    doc.add_paragraph("You should see Hello and the numpy version. That means Python and your libraries are working with Cursor.")
    doc.add_paragraph()

    # --- Quick reference ---
    add_heading(doc, "Quick reference", 1)
    doc.add_paragraph("Python — https://www.python.org/downloads — Check: python3 --version")
    doc.add_paragraph("pip — Comes with Python — Check: pip3 --version")
    doc.add_paragraph("Cursor — https://cursor.com — Open Cursor from Applications")
    doc.add_paragraph("Libraries — Via Terminal — pip3 install numpy pandas matplotlib jupyter requests")
    doc.add_paragraph()

    # --- Troubleshooting ---
    add_heading(doc, "Troubleshooting", 1)
    add_para(doc, '"python3: command not found"', bold=True)
    doc.add_paragraph("Python is not installed or not on your PATH. Install Python from python.org again, then close and reopen Terminal.")
    add_para(doc, '"pip3: command not found"', bold=True)
    doc.add_paragraph("Use python3 -m pip instead of pip3 in all the commands above (e.g. python3 -m pip install numpy).")
    add_para(doc, '"Permission denied" when installing libraries', bold=True)
    doc.add_paragraph("Do not use sudo. Install only for your user: pip3 install --user numpy (you can add --user to any pip3 install command if you get permission errors).")
    add_para(doc, "Cursor doesn't run or crashes", bold=True)
    doc.add_paragraph("Make sure you dragged Cursor into Applications. Try opening it again from Applications. If it still fails, download the installer again from cursor.com and reinstall.")
    doc.add_paragraph()
    add_para(doc, "That's it: IDE (Cursor), Python, and Python libraries only. No project folders or app setup.")

    out_path = __file__.replace("generate_setup_word_doc.py", "SETUP_IDE_PYTHON.docx")
    doc.save(out_path)
    print(f"Created: {out_path}")

if __name__ == "__main__":
    main()
